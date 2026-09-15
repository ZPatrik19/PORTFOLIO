"""Unit tests for DOCX parsing."""
from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from tkip.parsing import parse_document
from test_support.factories import make_document_record

pytestmark = pytest.mark.unit


def test_docx_parser_preserves_heading_paragraph_and_table(tmp_path: Path) -> None:
    # Arrange: create a tiny deterministic DOCX fixture.
    path = tmp_path / "sample.docx"
    word = Document()
    word.add_heading("Machine Learning", level=1)
    word.add_paragraph("Random forests aggregate multiple decision trees.")
    table = word.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Model"
    table.cell(0, 1).text = "Type"
    table.cell(1, 0).text = "Random Forest"
    table.cell(1, 1).text = "Ensemble"
    word.save(path)

    document = make_document_record(path, title="Sample", source_type="private")

    # Act
    blocks = parse_document(document)

    # Assert
    assert any(block.block_type == "heading" and "Machine Learning" in block.text for block in blocks)
    assert any("Random forests" in block.text for block in blocks)
    assert any(block.block_type == "table" and "Random Forest" in block.text for block in blocks)
