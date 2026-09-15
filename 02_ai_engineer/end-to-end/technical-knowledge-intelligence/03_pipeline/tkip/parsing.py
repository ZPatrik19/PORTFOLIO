"""Document parsers for PDF, DOCX, HTML, text/Markdown and optional EPUB files."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable

from bs4 import BeautifulSoup

from .config import PROJECT_ROOT, load_config
from .exceptions import DocumentParsingError
from .logging_config import get_logger
from .models import DocumentRecord, ParsedBlock

LOGGER = get_logger(__name__)

CODE_HINT = re.compile(
    r"(^\s{4,}|```|def |class |import |from .* import |SELECT |kubectl |docker |pip install|[{};])",
    re.MULTILINE | re.IGNORECASE,
)
HEADING_NUMBER_PATTERN = re.compile(r"^\d+(\.\d+)*\s+")
LIST_PATTERN = re.compile(r"^(\d+[.)]|[-*•])\s+")
CHAPTER_PATTERN = re.compile(r"^(chapter|chap\.|fejezet)\s+", re.IGNORECASE)
CAPTION_PATTERN = re.compile(r"^(figure|fig\.|table)\s*\d+", re.IGNORECASE)
EQUATION_CHARS = "=∑∫√λμσ→≤≥"


def classify_text_block(
    text: str,
    font_size: float | None = None,
    median_font_size: float | None = None,
) -> str:
    """Classify extracted text using lightweight structural heuristics.

    The classifier deliberately avoids model dependencies because it runs during
    ingestion and must remain deterministic and inexpensive for large libraries.
    """

    stripped = text.strip()
    if not stripped:
        return "unknown"
    if CODE_HINT.search(stripped):
        return "code"
    if CAPTION_PATTERN.match(stripped):
        return "caption"
    if sum(character in stripped for character in EQUATION_CHARS) >= 2 and len(stripped) < 500:
        return "equation"
    if LIST_PATTERN.match(stripped):
        return "list"
    if (
        font_size is not None
        and median_font_size is not None
        and font_size >= median_font_size * 1.25
        and len(stripped) < 180
    ):
        return "heading"
    if len(stripped) < 120 and (stripped.isupper() or HEADING_NUMBER_PATTERN.match(stripped)):
        return "heading"
    return "paragraph"


# Backward-compatible internal alias used by older tests/imports.
_classify = classify_text_block


def parse_document(document: DocumentRecord) -> list[ParsedBlock]:
    """Parse a supported document into logical blocks.

    Raises:
        DocumentParsingError: if the format is unsupported or parsing fails.
    """

    path = Path(document.path)
    extension = path.suffix.lower()
    parser_map = {
        ".pdf": _parse_pdf,
        ".html": _parse_html,
        ".htm": _parse_html,
        ".docx": _parse_docx,
        ".epub": _parse_epub,
        ".txt": _parse_text,
        ".md": _parse_text,
        ".markdown": _parse_text,
    }
    parser = parser_map.get(extension)
    if parser is None:
        raise DocumentParsingError(f"Unsupported document format '{extension}' for {path.name}")
    return parser(document, path)


def _parse_pdf(document: DocumentRecord, path: Path) -> list[ParsedBlock]:
    try:
        import pymupdf as fitz
    except ImportError as exc:
        raise DocumentParsingError("PDF parsing requires PyMuPDF (pymupdf).") from exc

    _configure_mupdf_diagnostics(fitz)
    parsing_config = load_config().get("parsing", {})
    extract_images = bool(parsing_config.get("extract_images", False))
    detect_tables = bool(parsing_config.get("detect_tables", False))
    blocks: list[ParsedBlock] = []
    current_heading: str | None = None
    current_chapter: str | None = None

    try:
        pdf = fitz.open(path)
        try:
            text_flags = _pdf_text_flags(fitz, extract_images)
            for page_number, page in enumerate(pdf, start=1):
                page_dict = page.get_text("dict", flags=text_flags)
                median_font_size = _median_font_size(page_dict)
                page_height = float(page.rect.height)

                for block_index, raw_block in enumerate(page_dict.get("blocks", [])):
                    if raw_block.get("type") == 1:
                        if extract_images and raw_block.get("image"):
                            blocks.append(
                                _extract_pdf_image(
                                    document,
                                    raw_block,
                                    page_number,
                                    block_index,
                                    current_heading,
                                )
                            )
                        continue
                    parsed = _parse_pdf_text_block(
                        document=document,
                        raw_block=raw_block,
                        page_number=page_number,
                        page_height=page_height,
                        median_font_size=median_font_size,
                        current_heading=current_heading,
                        current_chapter=current_chapter,
                    )
                    if parsed is None:
                        continue
                    block, current_heading, current_chapter = parsed
                    blocks.append(block)

                if detect_tables:
                    blocks.extend(
                        _extract_pdf_tables(
                            document,
                            page,
                            page_number,
                            current_heading,
                            current_chapter,
                        )
                    )
            document.page_count = len(pdf)
        finally:
            pdf.close()
    except Exception as exc:
        if isinstance(exc, DocumentParsingError):
            raise
        raise DocumentParsingError(f"PDF parsing failed for {path.name}: {exc}") from exc
    finally:
        _clear_mupdf_warnings(fitz)

    return blocks


def _configure_mupdf_diagnostics(fitz: object) -> None:
    """Suppress recoverable MuPDF console noise without hiding Python errors."""

    try:
        fitz.TOOLS.mupdf_display_errors(False)
        fitz.TOOLS.mupdf_display_warnings(False)
        fitz.TOOLS.reset_mupdf_warnings()
    except AttributeError:
        LOGGER.debug("MuPDF diagnostic controls are unavailable in this version")


def _clear_mupdf_warnings(fitz: object) -> None:
    try:
        fitz.TOOLS.mupdf_warnings(reset=True)
    except AttributeError:
        LOGGER.debug("MuPDF warning buffer reset is unavailable in this version")


def _pdf_text_flags(fitz: object, extract_images: bool) -> int:
    if extract_images:
        return fitz.TEXTFLAGS_DICT
    return fitz.TEXTFLAGS_DICT & ~fitz.TEXT_PRESERVE_IMAGES


def _median_font_size(page_dict: dict) -> float:
    sizes = [
        float(span.get("size", 10))
        for block in page_dict.get("blocks", [])
        for line in block.get("lines", [])
        for span in line.get("spans", [])
    ]
    if not sizes:
        return 10.0
    sizes.sort()
    return sizes[len(sizes) // 2]


def _parse_pdf_text_block(
    *,
    document: DocumentRecord,
    raw_block: dict,
    page_number: int,
    page_height: float,
    median_font_size: float,
    current_heading: str | None,
    current_chapter: str | None,
) -> tuple[ParsedBlock, str | None, str | None] | None:
    lines = raw_block.get("lines")
    if not lines:
        return None

    text = "\n".join(
        "".join(span.get("text", "") for span in line.get("spans", [])).strip()
        for line in lines
    ).strip()
    if not text:
        return None

    span_sizes = [
        float(span.get("size", 10))
        for line in lines
        for span in line.get("spans", [])
    ]
    average_size = sum(span_sizes) / len(span_sizes) if span_sizes else median_font_size
    block_type = classify_text_block(text, average_size, median_font_size)

    bbox = raw_block.get("bbox", (0, 0, 0, 0))
    y_start, y_end = float(bbox[1]), float(bbox[3])
    if len(text) < 180 and (y_end < page_height * 0.055 or y_start > page_height * 0.945):
        block_type = "header_footer"

    if block_type == "heading":
        current_heading = text[:240]
        if CHAPTER_PATTERN.match(text):
            current_chapter = text[:240]

    parsed = ParsedBlock(
        document_id=document.document_id,
        page=page_number,
        block_type=block_type,
        text=text,
        section=current_heading,
        chapter=current_chapter,
    )
    return parsed, current_heading, current_chapter


def _extract_pdf_image(
    document: DocumentRecord,
    raw_block: dict,
    page_number: int,
    block_index: int,
    current_heading: str | None,
) -> ParsedBlock:
    extension = raw_block.get("ext", "png")
    asset_dir = PROJECT_ROOT / "01_data" / "interim" / "figures" / document.document_id
    asset_dir.mkdir(parents=True, exist_ok=True)
    asset_path = asset_dir / f"page_{page_number:04d}_image_{block_index:03d}.{extension}"
    asset_path.write_bytes(raw_block["image"])
    return ParsedBlock(
        document_id=document.document_id,
        page=page_number,
        block_type="figure",
        text=f"[Extracted figure from page {page_number}]",
        section=current_heading,
        asset_path=str(asset_path.relative_to(PROJECT_ROOT)),
    )


def _extract_pdf_tables(
    document: DocumentRecord,
    page: object,
    page_number: int,
    current_heading: str | None,
    current_chapter: str | None,
) -> list[ParsedBlock]:
    try:
        found = page.find_tables()
    except (AttributeError, RuntimeError, ValueError) as exc:
        LOGGER.debug("Table detection skipped on page %s: %s", page_number, exc)
        return []

    blocks: list[ParsedBlock] = []
    for table in getattr(found, "tables", []) or []:
        data = table.extract()
        table_text = "\n".join(
            " | ".join("" if value is None else str(value) for value in row)
            for row in data
            if row
        ).strip()
        if not table_text:
            continue
        blocks.append(
            ParsedBlock(
                document_id=document.document_id,
                page=page_number,
                block_type="table",
                text=table_text,
                section=current_heading,
                chapter=current_chapter,
            )
        )
    return blocks


def _parse_html(document: DocumentRecord, path: Path) -> list[ParsedBlock]:
    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="ignore"), "html.parser")
    blocks: list[ParsedBlock] = []
    current_heading: str | None = None
    for node in soup.find_all(["h1", "h2", "h3", "h4", "p", "pre", "code", "li", "table"]):
        text = node.get_text(" ", strip=True)
        if not text:
            continue
        block_type = _html_block_type(node.name)
        if block_type == "heading":
            current_heading = text
        blocks.append(
            ParsedBlock(
                document_id=document.document_id,
                page=1,
                block_type=block_type,
                text=text,
                section=current_heading,
            )
        )
    document.page_count = 1
    return blocks


def _html_block_type(tag_name: str) -> str:
    if tag_name.startswith("h"):
        return "heading"
    if tag_name in {"pre", "code"}:
        return "code"
    if tag_name == "table":
        return "table"
    if tag_name == "li":
        return "list"
    return "paragraph"


def _parse_text(document: DocumentRecord, path: Path) -> list[ParsedBlock]:
    lines = path.read_text(encoding="utf-8", errors="ignore").splitlines()
    blocks: list[ParsedBlock] = []
    current_heading: str | None = None
    in_code_block = False
    buffer: list[str] = []

    def flush_buffer(block_type: str | None = None) -> None:
        if not buffer:
            return
        text = "\n".join(buffer).strip()
        buffer.clear()
        if not text:
            return
        blocks.append(
            ParsedBlock(
                document_id=document.document_id,
                page=1,
                block_type=block_type or classify_text_block(text),
                text=text,
                section=current_heading,
            )
        )

    for line in [*lines, ""]:
        stripped = line.strip()
        if stripped.startswith("```"):
            flush_buffer("code" if in_code_block else None)
            in_code_block = not in_code_block
            continue
        if in_code_block:
            buffer.append(line)
            continue
        if line.startswith("#"):
            flush_buffer()
            current_heading = line.lstrip("# ").strip()
            blocks.append(
                ParsedBlock(
                    document_id=document.document_id,
                    page=1,
                    block_type="heading",
                    text=current_heading,
                    section=current_heading,
                )
            )
        elif stripped:
            buffer.append(line)
        else:
            flush_buffer()

    document.page_count = 1
    return blocks


def _parse_docx(document: DocumentRecord, path: Path) -> list[ParsedBlock]:
    """Parse Word documents while preserving headings, paragraphs and tables."""

    try:
        from docx import Document
    except ImportError as exc:
        raise DocumentParsingError("DOCX support requires python-docx.") from exc

    try:
        word_document = Document(str(path))
        blocks: list[ParsedBlock] = []
        current_heading: str | None = None
        current_chapter: str | None = None

        for paragraph in word_document.paragraphs:
            text = (paragraph.text or "").strip()
            if not text:
                continue
            style = (getattr(getattr(paragraph, "style", None), "name", "") or "").lower()
            block_type = classify_text_block(text)
            if style.startswith("heading") or style in {"title", "subtitle"}:
                block_type = "heading"
                current_heading = text[:240]
                if style in {"title", "heading 1", "heading1"} or CHAPTER_PATTERN.match(text):
                    current_chapter = text[:240]
            elif style.startswith("list"):
                block_type = "list"

            blocks.append(
                ParsedBlock(
                    document_id=document.document_id,
                    page=1,
                    block_type=block_type,
                    text=text,
                    section=current_heading,
                    chapter=current_chapter,
                )
            )

        blocks.extend(
            _docx_table_blocks(document, word_document.tables, current_heading, current_chapter)
        )
        document.page_count = 1
        return blocks
    except Exception as exc:
        if isinstance(exc, DocumentParsingError):
            raise
        raise DocumentParsingError(f"DOCX parsing failed for {path.name}: {exc}") from exc


def _docx_table_blocks(
    document: DocumentRecord,
    tables: Iterable[object],
    current_heading: str | None,
    current_chapter: str | None,
) -> list[ParsedBlock]:
    blocks: list[ParsedBlock] = []
    for table in tables:
        rows: list[str] = []
        for row in table.rows:
            values = [(cell.text or "").strip().replace("\n", " ") for cell in row.cells]
            if any(values):
                rows.append(" | ".join(values))
        table_text = "\n".join(rows).strip()
        if table_text:
            blocks.append(
                ParsedBlock(
                    document_id=document.document_id,
                    page=1,
                    block_type="table",
                    text=table_text,
                    section=current_heading,
                    chapter=current_chapter,
                )
            )
    return blocks


def _parse_epub(document: DocumentRecord, path: Path) -> list[ParsedBlock]:
    try:
        from ebooklib import ITEM_DOCUMENT, epub
    except ImportError as exc:
        raise DocumentParsingError(
            "EPUB support requires ebooklib; install it or convert EPUB to PDF/HTML."
        ) from exc

    try:
        book = epub.read_epub(str(path))
        blocks: list[ParsedBlock] = []
        page_number = 0
        for item in book.get_items_of_type(ITEM_DOCUMENT):
            page_number += 1
            soup = BeautifulSoup(item.get_content(), "html.parser")
            for node in soup.find_all(["h1", "h2", "h3", "p", "pre", "code", "li"]):
                text = node.get_text(" ", strip=True)
                if not text:
                    continue
                blocks.append(
                    ParsedBlock(
                        document_id=document.document_id,
                        page=page_number,
                        block_type=_html_block_type(node.name),
                        text=text,
                    )
                )
        document.page_count = page_number
        return blocks
    except Exception as exc:
        if isinstance(exc, DocumentParsingError):
            raise
        raise DocumentParsingError(f"EPUB parsing failed for {path.name}: {exc}") from exc
