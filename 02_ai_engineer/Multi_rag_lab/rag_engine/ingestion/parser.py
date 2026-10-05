from __future__ import annotations

import hashlib
import json
from pathlib import Path

from bs4 import BeautifulSoup

from rag_engine.models import Document


class ParseError(RuntimeError):
    pass


def _document_id(path: Path, suffix: str = "") -> str:
    raw = f"{path.resolve()}::{suffix}".encode()
    return hashlib.sha256(raw).hexdigest()[:20]



def _base_metadata(path: Path) -> dict[str, object]:
    metadata: dict[str, object] = {"source": str(path), "title": path.name}
    sidecar = path.with_suffix(path.suffix + ".source.json")
    if sidecar.exists():
        try:
            source_info = json.loads(sidecar.read_text(encoding="utf-8"))
            metadata["local_path"] = str(path)
            metadata["source"] = source_info.get("url", str(path))
            metadata["sha256"] = source_info.get("sha256")
            for key in ("source_id", "title", "category", "organization", "language", "format", "source_page", "description"):
                if source_info.get(key) not in (None, ""):
                    metadata[key] = source_info[key]
        except (OSError, json.JSONDecodeError):
            pass
    return metadata


def _read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def parse_file(path: Path) -> list[Document]:
    if not path.exists():
        raise ParseError(f"File does not exist: {path}")
    ext = path.suffix.lower()
    if ext in {".txt", ".md", ".markdown"}:
        text = _read_text(path)
        return [Document(document_id=_document_id(path), text=text, metadata=_base_metadata(path))]
    if ext in {".html", ".htm"}:
        soup = BeautifulSoup(_read_text(path), "html.parser")
        title = soup.title.string.strip() if soup.title and soup.title.string else path.name
        # Website chrome is retrieval noise, especially when tens/hundreds of pages from the
        # same domain repeat navigation, quick links and footers. Keep the article/main body.
        for tag in soup(["script", "style", "noscript", "nav", "footer", "header", "form", "aside"]):
            tag.decompose()
        content_root = soup.find("main") or soup.find("article") or soup.body or soup
        # Preserve H1-H3 as Markdown-style boundaries so structure-aware chunking
        # can operate on HTML input after parsing instead of losing hierarchy.
        for level in (1, 2, 3):
            for heading in content_root.find_all(f"h{level}"):
                heading_text = heading.get_text(" ", strip=True)
                heading.clear()
                heading.append(f"{'#' * level} {heading_text}")
        text = content_root.get_text("\n")
        return [Document(document_id=_document_id(path), text=text, metadata={**_base_metadata(path), "title": title})]
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise ParseError("Install pypdf to parse PDF files") from exc
        try:
            reader = PdfReader(path)
            docs = []
            for index, page in enumerate(reader.pages, start=1):
                text = page.extract_text() or ""
                docs.append(Document(
                    document_id=_document_id(path, str(index)),
                    text=text,
                    metadata={**_base_metadata(path), "page": index},
                ))
            return docs
        except Exception as exc:
            raise ParseError(f"Could not parse PDF {path}: {exc}") from exc
    if ext == ".docx":
        try:
            from docx import Document as DocxDocument
        except ImportError as exc:
            raise ParseError("Install python-docx to parse DOCX files") from exc
        docx = DocxDocument(path)
        text = "\n".join(p.text for p in docx.paragraphs)
        return [Document(document_id=_document_id(path), text=text, metadata=_base_metadata(path))]
    raise ParseError(f"Unsupported document type: {ext}")
